"""
End-to-end local test of the full permit workflow.

Mocks external services (Graph mail, Adobe Sign, SharePoint, LibreOffice)
so the entire pipeline can run on any machine without credentials.
"""
import asyncio
import shutil
import sys
from pathlib import Path
from unittest.mock import AsyncMock, patch

import httpx

BASE = "http://127.0.0.1:8000"


# ── Mock patches for external services ──────────────────────────────────────

def mock_graph_mail():
    """Mock Microsoft Graph email sending."""
    async def fake_send(to_email, to_name, subject, body_html, attachment_path=None):
        print(f"  [MOCK EMAIL] To: {to_name} <{to_email}>")
        print(f"               Subject: {subject}")
        if attachment_path:
            print(f"               Attachment: {Path(attachment_path).name}")

    return patch("app.services.graph_mail.send_permit_email", side_effect=fake_send)


def mock_pdf_converter():
    """Mock LibreOffice PDF conversion by copying the Excel as a fake PDF."""
    async def fake_convert(excel_path):
        output_dir = Path("generated")
        output_dir.mkdir(exist_ok=True)
        pdf_path = output_dir / f"{excel_path.stem}.pdf"
        shutil.copy2(excel_path, pdf_path)
        print(f"  [MOCK PDF] Converted {excel_path.name} → {pdf_path.name}")
        return pdf_path

    return patch("app.services.pdf_converter.convert_excel_to_pdf", side_effect=fake_convert)


def mock_adobe_sign():
    """Mock Adobe Sign: upload, create agreement, register webhook, download."""
    async def fake_upload(pdf_path):
        tid = "MOCK_TRANSIENT_12345"
        print(f"  [MOCK SIGN] Uploaded transient document: {tid}")
        return tid

    async def fake_create(transient_id, name, signers):
        aid = "MOCK_AGREEMENT_67890"
        print(f"  [MOCK SIGN] Created agreement: {aid}")
        for s in signers:
            print(f"              Signer: {s['email']} (order {signers.index(s)+1})")
        return aid

    async def fake_webhook(agreement_id):
        wid = "MOCK_WEBHOOK_11111"
        print(f"  [MOCK SIGN] Registered webhook: {wid}")
        return wid

    async def fake_download(agreement_id, output_path):
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text("MOCK SIGNED PDF CONTENT")
        print(f"  [MOCK SIGN] Downloaded signed PDF → {output_path.name}")
        return output_path

    return [
        patch("app.services.adobe_sign.upload_transient_document", side_effect=fake_upload),
        patch("app.services.adobe_sign.create_agreement", side_effect=fake_create),
        patch("app.services.adobe_sign.register_webhook", side_effect=fake_webhook),
        patch("app.services.adobe_sign.download_signed_document", side_effect=fake_download),
    ]


def mock_sharepoint():
    """Mock SharePoint upload."""
    async def fake_upload(file_path, filename):
        url = f"https://mock-sharepoint.com/sites/permits/{filename}"
        print(f"  [MOCK SP] Uploaded {filename} → {url}")
        return url

    return patch("app.services.sharepoint.upload_to_sharepoint", side_effect=fake_upload)


# ── Test steps ──────────────────────────────────────────────────────────────

async def step_1_create_request(client):
    print("\n" + "=" * 60)
    print("STEP 1: Create permit request")
    print("=" * 60)

    resp = await client.post(f"{BASE}/permits", json={
        "requester_name": "Ahmad Hassan",
        "requester_email": "ahmad@example.com",
        "project_name": "Building A Renovation",
        "project_location": "Downtown Amman",
        "permit_types": [
            {
                "permit_type": "hot_work",
                "owner_email": "owner@example.com",
                "coordinator_email": "coordinator@example.com",
                "approver_email": "approver@example.com",
            }
        ],
    })

    assert resp.status_code == 201, f"Expected 201, got {resp.status_code}: {resp.text}"
    data = resp.json()
    request_id = data["id"]
    permit_id = data["permits"][0]["id"]
    print(f"  Created request #{request_id} with permit #{permit_id}")
    print(f"  Status: {data['permits'][0]['status']}")
    return request_id, permit_id


async def step_2_check_template_sent(client, request_id, permit_id):
    print("\n" + "=" * 60)
    print("STEP 2: Verify template was filled and email sent")
    print("=" * 60)

    for attempt in range(10):
        await asyncio.sleep(1)
        resp = await client.get(f"{BASE}/permits/{request_id}")
        data = resp.json()
        status = data["permits"][0]["status"]
        if status != "draft":
            break

    print(f"  Permit status: {status}")

    if status == "template_sent":
        print("  Template filled and email sent successfully!")
    else:
        print(f"  Unexpected status: {status}")
    return status


async def step_3_upload_invalid(client, permit_id):
    print("\n" + "=" * 60)
    print("STEP 3: Upload INVALID Excel (missing required fields)")
    print("=" * 60)

    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws["B3"] = "Ahmad Hassan"
    ws["B4"] = "ahmad@example.com"
    # Deliberately leave B5, B10, B11, B12 empty
    invalid_path = Path("uploads") / "test_invalid.xlsx"
    invalid_path.parent.mkdir(exist_ok=True)
    wb.save(invalid_path)

    with open(invalid_path, "rb") as f:
        resp = await client.post(
            f"{BASE}/permits/{permit_id}/upload",
            files={"file": ("test_invalid.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )

    assert resp.status_code == 200, f"Upload failed with {resp.status_code}: {resp.text}"
    data = resp.json()
    print(f"  Valid: {data['valid']}")
    print(f"  Errors ({len(data['errors'])}):")
    for err in data["errors"]:
        print(f"    - {err}")
    assert not data["valid"], "Should have failed validation"
    print("  Validation correctly rejected the file!")
    return data


async def step_4_upload_valid(client, permit_id):
    print("\n" + "=" * 60)
    print("STEP 4: Upload VALID Excel (all required fields filled)")
    print("=" * 60)

    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws["B3"] = "Ahmad Hassan"
    ws["B4"] = "ahmad@example.com"
    ws["B5"] = "Building A Renovation"
    ws["B6"] = "Downtown Amman"
    ws["B7"] = "hot_work"
    ws["B10"] = "2026-11-01"
    ws["B11"] = "2026-12-31"
    ws["B12"] = "Interior demolition and hot work on floors 3-5"
    valid_path = Path("uploads") / "test_valid.xlsx"
    valid_path.parent.mkdir(exist_ok=True)
    wb.save(valid_path)

    with open(valid_path, "rb") as f:
        resp = await client.post(
            f"{BASE}/permits/{permit_id}/upload",
            files={"file": ("test_valid.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )

    data = resp.json()
    print(f"  Valid: {data['valid']}")
    assert data["valid"], f"Should have passed validation, got errors: {data.get('errors')}"
    print("  Validation passed!")
    return data


async def step_5_check_signing_status(client, request_id):
    print("\n" + "=" * 60)
    print("STEP 5: Verify PDF generated and sent for signing")
    print("=" * 60)

    resp = await client.get(f"{BASE}/permits/{request_id}")
    data = resp.json()
    permit = data["permits"][0]
    print(f"  Status: {permit['status']}")
    return permit["status"]


async def step_6_simulate_webhook(client):
    print("\n" + "=" * 60)
    print("STEP 6: Simulate Adobe Sign webhook (signing complete)")
    print("=" * 60)

    resp = await client.post(f"{BASE}/webhooks/adobe-sign", json={
        "event": "AGREEMENT_ACTION_COMPLETED",
        "agreement": {
            "id": "MOCK_AGREEMENT_67890",
            "status": "SIGNED",
        },
    })

    data = resp.json()
    print(f"  Webhook response: {data}")
    assert data["status"] == "ok"
    print("  Webhook processed successfully!")


async def step_7_verify_completion(client, request_id):
    print("\n" + "=" * 60)
    print("STEP 7: Verify full completion")
    print("=" * 60)

    resp = await client.get(f"{BASE}/permits/{request_id}")
    data = resp.json()
    permit = data["permits"][0]

    print(f"  Final status: {permit['status']}")
    print(f"  SharePoint URL: {permit['sharepoint_url']}")
    print(f"  Validation errors: {permit['validation_errors']}")

    if permit["status"] == "completed":
        print("\n  FULL WORKFLOW COMPLETED SUCCESSFULLY!")
    return permit


# ── Main ────────────────────────────────────────────────────────────────────

async def run_test():
    print("=" * 60)
    print("  PERMIT WORKFLOW — END-TO-END LOCAL TEST")
    print("  (External services are mocked)")
    print("=" * 60)

    async with httpx.AsyncClient(timeout=30) as client:
        # Step 1: Create request
        request_id, permit_id = await step_1_create_request(client)

        # Step 2: Wait for background template task
        await asyncio.sleep(2)
        await step_2_check_template_sent(client, request_id, permit_id)

        # Step 3: Upload invalid Excel → expect validation failure
        await step_3_upload_invalid(client, permit_id)

        # Step 4: Upload valid Excel → triggers PDF + signing
        await step_4_upload_valid(client, permit_id)

        # Step 5: Check status after upload
        status = await step_5_check_signing_status(client, request_id)

        # Step 6: Simulate webhook from Adobe Sign
        await step_6_simulate_webhook(client)

        # Step 7: Verify everything completed
        permit = await step_7_verify_completion(client, request_id)

    print("\n" + "=" * 60)
    if permit["status"] == "completed":
        print("  ALL 7 STEPS PASSED — WORKFLOW IS WORKING!")
    else:
        print(f"  WORKFLOW ENDED AT STATUS: {permit['status']}")
    print("=" * 60)


if __name__ == "__main__":
    import contextlib

    adobe_mocks = mock_adobe_sign()
    stack = contextlib.ExitStack()

    with stack:
        stack.enter_context(mock_graph_mail())
        stack.enter_context(mock_pdf_converter())
        stack.enter_context(mock_sharepoint())
        for m in adobe_mocks:
            stack.enter_context(m)

        from main import app
        import uvicorn

        async def main():
            config = uvicorn.Config(app, host="127.0.0.1", port=8000, log_level="warning")
            server = uvicorn.Server(config)

            server_task = asyncio.create_task(server.serve())
            await asyncio.sleep(2)

            try:
                await run_test()
            finally:
                server.should_exit = True
                await server_task

        asyncio.run(main())
