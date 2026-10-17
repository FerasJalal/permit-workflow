"""Run once to generate sample Excel templates in the templates/ directory."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pathlib import Path


def create_template(permit_type: str, output_dir: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Permit Application"

    header_font = Font(bold=True, size=14, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    label_font = Font(bold=True, size=11)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 45

    ws.merge_cells("A1:B1")
    ws["A1"] = f"{permit_type} Permit Application"
    ws["A1"].font = header_font
    ws["A1"].fill = header_fill
    ws["A1"].alignment = Alignment(horizontal="center")

    fields = [
        (3, "Requester Name"),
        (4, "Requester Email"),
        (5, "Project Name"),
        (6, "Project Location"),
        (7, "Permit Type"),
        (8, ""),
        (9, "--- Fill the fields below ---"),
        (10, "Start Date"),
        (11, "End Date"),
        (12, "Description of Work"),
        (13, "Equipment Used"),
        (14, "Safety Measures"),
        (15, "Additional Notes"),
    ]

    for row, label in fields:
        if label.startswith("---"):
            ws.merge_cells(f"A{row}:B{row}")
            ws[f"A{row}"] = label.strip("- ")
            ws[f"A{row}"].font = Font(bold=True, italic=True, size=11)
            ws[f"A{row}"].alignment = Alignment(horizontal="center")
            continue
        if not label:
            continue
        ws[f"A{row}"] = label
        ws[f"A{row}"].font = label_font
        ws[f"A{row}"].border = thin_border
        ws[f"B{row}"].border = thin_border

    output_path = output_dir / f"{permit_type}.xlsx"
    wb.save(output_path)
    print(f"Created template: {output_path}")


if __name__ == "__main__":
    templates_dir = Path("templates")
    templates_dir.mkdir(exist_ok=True)

    for ptype in ["hot_work", "confined_space", "excavation", "electrical", "general"]:
        create_template(ptype, templates_dir)

    print("Done! Templates created in templates/")
