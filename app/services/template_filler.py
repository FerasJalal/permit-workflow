import shutil
from pathlib import Path

from openpyxl import load_workbook

from app.config import settings


FIELD_MAPPING = {
    "requester_name": "B3",
    "requester_email": "B4",
    "project_name": "B5",
    "project_location": "B6",
    "permit_type": "B7",
}


async def fill_template(permit_type: str, data: dict) -> Path:
    template_path = settings.templates_path / f"{permit_type}.xlsx"
    if not template_path.exists():
        raise FileNotFoundError(f"Template not found: {template_path}")

    output_dir = settings.uploads_path / str(data.get("permit_id", "unknown"))
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"{permit_type}_filled.xlsx"

    shutil.copy2(template_path, output_path)

    wb = load_workbook(output_path)
    ws = wb.active

    for field, cell in FIELD_MAPPING.items():
        if field in data:
            ws[cell] = data[field]

    wb.save(output_path)
    wb.close()

    return output_path
