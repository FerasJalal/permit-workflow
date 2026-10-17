from pathlib import Path

from openpyxl import load_workbook

from app.schemas import ValidationResult

REQUIRED_CELLS = {
    "B3": "Requester Name",
    "B4": "Requester Email",
    "B5": "Project Name",
    "B6": "Project Location",
    "B7": "Permit Type",
    "B10": "Start Date",
    "B11": "End Date",
    "B12": "Description of Work",
}


async def validate_excel(file_path: Path) -> ValidationResult:
    wb = load_workbook(file_path, data_only=True)
    ws = wb.active

    errors = []
    for cell, label in REQUIRED_CELLS.items():
        value = ws[cell].value
        if value is None or (isinstance(value, str) and value.strip() == ""):
            errors.append(f"Required field '{label}' (cell {cell}) is empty")

    wb.close()

    return ValidationResult(valid=len(errors) == 0, errors=errors)
